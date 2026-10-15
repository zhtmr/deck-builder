#!/usr/bin/env python3
"""
export_pptx.py — 장표 HTML 을 슬라이드별 스크린샷 + 발표자 노트가 들어간 PPTX 로 내보낸다.

애니메이션은 PPTX 로 옮겨지지 않는다. 각 슬라이드의 "모든 스텝이 끝난 최종 상태"를 이미지로 넣으므로
제출·공유용이다 (모션이 필요한 발표는 HTML 로 직접). 텍스트 편집은 불가능한 이미지 슬라이드다.

필요: pip install playwright python-pptx && python -m playwright install chromium
사용: python3 export_pptx.py out/deck.html -o out/deck.pptx
"""
import argparse
import sys
import tempfile
from pathlib import Path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("html", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=None)
    a = ap.parse_args()
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        sys.exit("playwright 가 없습니다: pip install playwright && python -m playwright install chromium")
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx 가 없습니다: pip install python-pptx")

    out = a.out or a.html.with_suffix(".pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9, 13.333in x 7.5in
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={"width": 1600, "height": 900}, device_scale_factor=1.5)
        pg.goto(a.html.resolve().as_uri() + "?still=1")
        pg.wait_for_timeout(500)
        n = pg.evaluate("deck.count")
        for i in range(n):
            pg.evaluate(f"deck.go({i}, deck.steps({i}), true)")
            pg.wait_for_timeout(250)
            img = Path(tmp) / f"s{i:02d}.png"
            pg.screenshot(path=str(img))
            notes = pg.evaluate("document.querySelector('.slide.active .notes')?.textContent || ''").strip()
            sl = prs.slides.add_slide(blank)
            sl.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
            if notes:
                sl.notes_slide.notes_text_frame.text = notes
        b.close()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"생성 완료: {out}  (슬라이드 {n}장, 이미지 + 발표자 노트)")


if __name__ == "__main__":
    main()
